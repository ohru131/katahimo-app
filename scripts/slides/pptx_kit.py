# -*- coding: utf-8 -*-
"""doc/12・doc/13 のPowerPoint資料を生成するための共通部品。

図形ベースで描くので、生成後もPowerPoint側で自由に編集できる
(画像を貼り込むのではなく、すべてネイティブの図形・テキスト・表として出力する)。
"""
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

# 配色(印刷しても読める明るいテーマ)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x14, 0x22, 0x3A)
MUTED = RGBColor(0x64, 0x74, 0x8B)
LINE = RGBColor(0xCB, 0xD5, 0xE1)
CARD = RGBColor(0xF1, 0xF5, 0xF9)
CARD2 = RGBColor(0xE2, 0xE8, 0xF0)
ACCENT = RGBColor(0x03, 0x69, 0xA1)
ACCENT_L = RGBColor(0xE0, 0xF2, 0xFE)
VIOLET = RGBColor(0x6D, 0x28, 0xD9)
VIOLET_L = RGBColor(0xED, 0xE9, 0xFE)
GREEN = RGBColor(0x04, 0x78, 0x57)
GREEN_L = RGBColor(0xD1, 0xFA, 0xE5)
AMBER = RGBColor(0xB4, 0x53, 0x09)
AMBER_L = RGBColor(0xFE, 0xF3, 0xC7)
RED = RGBColor(0xB9, 0x1C, 0x1C)
RED_L = RGBColor(0xFE, 0xE2, 0xE2)
PINK = RGBColor(0xBE, 0x18, 0x5D)
PINK_L = RGBColor(0xFC, 0xE7, 0xF3)
NAVY = RGBColor(0x0F, 0x17, 0x2A)
ORANGE = RGBColor(0xC2, 0x41, 0x0C)
ORANGE_L = RGBColor(0xFF, 0xED, 0xD5)
TEAL = RGBColor(0x0F, 0x76, 0x6E)
TEAL_L = RGBColor(0xCC, 0xFB, 0xF1)

FONT = "Meiryo"
MONO = "Consolas"

SLIDE_W = 13.333
SLIDE_H = 7.5
ML = 0.5
MR = 0.5
CW = SLIDE_W - ML - MR
BODY_TOP = 1.18
BODY_BOTTOM = 6.92


def _font(run, name=FONT):
    """日本語が別フォントに化けないよう latin/ea/cs すべてを指定する。"""
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    latin = rPr.find(qn("a:latin"))
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {"typeface": name}, None)
            if latin is not None:
                latin.addnext(el)
            else:
                rPr.append(el)
        else:
            el.set("typeface", name)


def _set_bullet(p, char="▪", font=FONT, color=None):
    pPr = p._p.get_or_add_pPr()
    if color is not None:
        buClr = pPr.makeelement(qn("a:buClr"), {}, None)
        srgb = pPr.makeelement(qn("a:srgbClr"), {"val": str(color)}, None)
        buClr.append(srgb)
        pPr.append(buClr)
    pPr.append(pPr.makeelement(qn("a:buFont"), {"typeface": font}, None))
    pPr.append(pPr.makeelement(qn("a:buChar"), {"char": char}, None))


def _no_bullet(p):
    pPr = p._p.get_or_add_pPr()
    pPr.append(pPr.makeelement(qn("a:buNone"), {}, None))


def _indent(p, mar_in, hang_in=0.0):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(int(Inches(mar_in))))
    pPr.set("indent", str(int(Inches(-hang_in))))


def new_deck():
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    return prs


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def rect(sl, x, y, w, h, fill=None, border=None, border_w=1.0,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12, dash=False):
    sh = sl.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            sh.adjustments[0] = radius
        except Exception:
            pass
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if border is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = border
        sh.line.width = Pt(border_w)
        if dash:
            sh.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    sh.shadow.inherit = False  # 既定の影を消す(空の effectLst が入る)
    sh.text_frame.word_wrap = True
    return sh


def text(sl, x, y, w, h, content, size=13, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, font=FONT, line=1.25, italic=False, space_after=0):
    """content は str か [(文字列, {size/color/bold/font}), ...]。"\\n" 単独で改段落。"""
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    _no_bullet(p)
    if line:
        p.line_spacing = line
    if space_after:
        p.space_after = Pt(space_after)
    runs = content if isinstance(content, list) else [(content, {})]
    for s, opt in runs:
        if s == "\n":
            p = tf.add_paragraph()
            p.alignment = align
            _no_bullet(p)
            if line:
                p.line_spacing = line
            if space_after:
                p.space_after = Pt(space_after)
            continue
        r = p.add_run()
        r.text = s
        r.font.size = Pt(opt.get("size", size))
        r.font.bold = opt.get("bold", bold)
        r.font.italic = opt.get("italic", italic)
        r.font.color.rgb = opt.get("color", color)
        _font(r, opt.get("font", font))
    return tb


def fill_text(sh, content, size=13, color=INK, bold=False, align=PP_ALIGN.CENTER,
              anchor=MSO_ANCHOR.MIDDLE, font=FONT, line=1.16, pad=0.045):
    tf = sh.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(pad)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    _no_bullet(p)
    p.line_spacing = line
    runs = content if isinstance(content, list) else [(content, {})]
    for s, opt in runs:
        if s == "\n":
            p = tf.add_paragraph()
            p.alignment = align
            _no_bullet(p)
            p.line_spacing = line
            continue
        r = p.add_run()
        r.text = s
        r.font.size = Pt(opt.get("size", size))
        r.font.bold = opt.get("bold", bold)
        r.font.color.rgb = opt.get("color", color)
        _font(r, opt.get("font", font))
    return sh


def box(sl, x, y, w, h, content, fill=CARD, border=LINE, color=INK, size=12, bold=False,
        border_w=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE, align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE, radius=0.14, dash=False, line=1.16):
    sh = rect(sl, x, y, w, h, fill=fill, border=border, border_w=border_w, shape=shape,
              radius=radius, dash=dash)
    fill_text(sh, content, size=size, color=color, bold=bold, align=align, anchor=anchor,
              line=line)
    return sh


def arrow(sl, p1, p2, color=MUTED, width=1.5, dash=False, head="triangle", tail=None,
          elbow=False):
    kind = MSO_CONNECTOR.ELBOW if elbow else MSO_CONNECTOR.STRAIGHT
    con = sl.shapes.add_connector(kind, Inches(p1[0]), Inches(p1[1]),
                                  Inches(p2[0]), Inches(p2[1]))
    con.shadow.inherit = False
    con.line.color.rgb = color
    con.line.width = Pt(width)
    if dash:
        con.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    ln = con.line._get_or_add_ln()
    if tail:
        ln.append(ln.makeelement(qn("a:headEnd"), {"type": tail, "w": "med", "len": "med"}, None))
    if head:
        ln.append(ln.makeelement(qn("a:tailEnd"), {"type": head, "w": "med", "len": "med"}, None))
    return con


def hline(sl, x1, y, x2, color=LINE, width=1.0, dash=False):
    return arrow(sl, (x1, y), (x2, y), color=color, width=width, dash=dash, head=None)


def vline(sl, x, y1, y2, color=LINE, width=1.0, dash=False):
    return arrow(sl, (x, y1), (x, y2), color=color, width=width, dash=dash, head=None)


def bullets(sl, x, y, w, h, items, size=12.5, color=INK, line=1.3, gap=4.5,
            marker="▪", marker_color=None):
    """items: [{'t': 本文 or run list, 'lv': 0..2, 'c': 色, 'b': 太字, 's': サイズ, 'mk': 記号}]"""
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        lv = it.get("lv", 0)
        mk = it.get("mk", marker if lv == 0 else "–")
        if mk:
            _set_bullet(p, mk, color=it.get("mkc", marker_color or (ACCENT if lv == 0 else MUTED)))
            _indent(p, 0.2 + 0.24 * lv, 0.2)
        else:
            _no_bullet(p)
            _indent(p, 0.24 * lv, 0)
        p.line_spacing = it.get("line", line)
        p.space_before = Pt(0 if i == 0 else it.get("gap", gap))
        runs = it["t"] if isinstance(it["t"], list) else [(it["t"], {})]
        for s, opt in runs:
            r = p.add_run()
            r.text = s
            r.font.size = Pt(opt.get("size", it.get("s", size)))
            r.font.bold = opt.get("bold", it.get("b", False))
            r.font.color.rgb = opt.get("color", it.get("c", color))
            _font(r, opt.get("font", FONT))
    return tb


def title_slide(prs, eyebrow, title, subtitle, meta):
    sl = blank(prs)
    rect(sl, 0, 0, SLIDE_W, SLIDE_H, fill=NAVY, border=None, shape=MSO_SHAPE.RECTANGLE)
    rect(sl, 0, 0, 0.22, SLIDE_H, fill=ACCENT, border=None, shape=MSO_SHAPE.RECTANGLE)
    text(sl, 1.1, 2.0, 11, 0.4, eyebrow, size=13.5, color=RGBColor(0x7D, 0xD3, 0xFC), bold=True)
    text(sl, 1.1, 2.45, 11.2, 1.6, title, size=38, color=WHITE, bold=True, line=1.18)
    text(sl, 1.1, 4.35, 11, 0.5, subtitle, size=16, color=RGBColor(0xCB, 0xD5, 0xE1), line=1.4)
    hline(sl, 1.1, 5.12, 5.2, color=RGBColor(0x47, 0x55, 0x69), width=1.2)
    text(sl, 1.1, 5.32, 11, 1.0, meta, size=12.5, color=RGBColor(0x94, 0xA3, 0xB8), line=1.55)
    return sl


def section_slide(prs, no, title, subtitle=""):
    sl = blank(prs)
    rect(sl, 0, 0, SLIDE_W, SLIDE_H, fill=CARD, border=None, shape=MSO_SHAPE.RECTANGLE)
    rect(sl, 0, 2.55, SLIDE_W, 2.4, fill=WHITE, border=None, shape=MSO_SHAPE.RECTANGLE)
    rect(sl, 0.9, 2.55, 0.1, 2.4, fill=ACCENT, border=None, shape=MSO_SHAPE.RECTANGLE)
    text(sl, 1.35, 2.95, 10, 0.5, no, size=14, color=ACCENT, bold=True)
    text(sl, 1.35, 3.38, 11, 0.9, title, size=31, color=INK, bold=True, line=1.2)
    if subtitle:
        text(sl, 1.35, 4.3, 11, 0.5, subtitle, size=14, color=MUTED, line=1.4)
    return sl


def slide(prs, title, small="", page=None, source="", accent=ACCENT):
    sl = blank(prs)
    rect(sl, 0, 0, SLIDE_W, 0.055, fill=accent, border=None, shape=MSO_SHAPE.RECTANGLE)
    runs = [(title, {"size": 22.5, "bold": True, "color": INK})]
    if small:
        runs.append(("   " + small, {"size": 12.5, "color": MUTED, "bold": False}))
    text(sl, ML, 0.36, CW, 0.57, runs, anchor=MSO_ANCHOR.BOTTOM)
    hline(sl, ML, 1.03, SLIDE_W - MR, color=LINE, width=1.4)
    if source:
        text(sl, ML, 7.02, CW - 0.6, 0.3, source, size=9, color=MUTED)
    if page is not None:
        text(sl, SLIDE_W - MR - 0.7, 7.02, 0.7, 0.3, str(page), size=10, color=MUTED,
             align=PP_ALIGN.RIGHT)
    return sl


def card(sl, x, y, w, h, title, items=None, accent=ACCENT, body_size=11.5, title_size=12.5,
         fill=WHITE, border=LINE, pad=0.18, body=None, line=1.32, gap=4):
    rect(sl, x, y, w, h, fill=fill, border=border, border_w=1.0)
    rect(sl, x, y, 0.075, h, fill=accent, border=None, shape=MSO_SHAPE.RECTANGLE)
    text(sl, x + pad + 0.05, y + 0.12, w - pad * 2, 0.32, title, size=title_size, color=accent,
         bold=True)
    top = y + 0.12 + 0.33
    if items:
        bullets(sl, x + pad + 0.05, top, w - pad * 2 - 0.05, h - (top - y) - 0.1, items,
                size=body_size, line=line, gap=gap)
    elif body:
        text(sl, x + pad + 0.05, top, w - pad * 2 - 0.05, h - (top - y) - 0.1, body,
             size=body_size, line=line)
    return top


def note(sl, x, y, w, h, label, body, accent=AMBER, fill=AMBER_L, size=11.5, lsize=10.5):
    """指定した h では本文がはみ出す場合、必要な高さまで自動で広げる。
    下端(6.98in)を越える場合は上へずらす。"""
    avail = w - 0.42
    lines = max(1, int(_text_w(body, size) / avail) + 1)
    need = (0.36 if label else 0.12) + lines * (size * 1.34 / 72) + 0.12
    h = max(h, need)
    if y + h > 6.98:
        y = 6.98 - h
    rect(sl, x, y, w, h, fill=fill, border=None)
    rect(sl, x, y, 0.075, h, fill=accent, border=None, shape=MSO_SHAPE.RECTANGLE)
    if label:
        text(sl, x + 0.2, y + 0.1, w - 0.4, 0.26, label, size=lsize, color=accent, bold=True)
    text(sl, x + 0.2, y + (0.38 if label else 0.12), w - 0.38, h - 0.46, body, size=size,
         color=INK, line=1.33)


def badge(sl, x, y, w, h, label, color=ACCENT, fill=ACCENT_L, size=10):
    sh = rect(sl, x, y, w, h, fill=fill, border=None, radius=0.5)
    fill_text(sh, label, size=size, color=color, bold=True)
    return sh


def table(sl, x, y, w, headers, rows, col_w=None, size=10.5, hsize=10.5, row_h=0.3,
          header_h=0.34, head_fill=CARD2, cell_colors=None, first_bold=False,
          aligns=None):
    nr, nc = len(rows) + 1, len(headers)
    gf = sl.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w),
                             Inches(header_h + row_h * len(rows)))
    tbl = gf.table
    tbl.first_row = True
    if col_w:
        total = sum(col_w)
        for i, cw in enumerate(col_w):
            tbl.columns[i].width = Emu(int(Inches(w) * cw / total))
    tbl.rows[0].height = Inches(header_h)
    for i in range(len(rows)):
        tbl.rows[i + 1].height = Inches(row_h)

    def put(cell, s, bold=False, color=INK, sz=size, align=PP_ALIGN.LEFT):
        cell.margin_left = cell.margin_right = Inches(0.07)
        cell.margin_top = cell.margin_bottom = Inches(0.03)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf = cell.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = align
        _no_bullet(p)
        p.line_spacing = 1.15
        r = p.add_run()
        r.text = s
        r.font.size = Pt(sz)
        r.font.bold = bold
        r.font.color.rgb = color
        _font(r)

    for c, htxt in enumerate(headers):
        cell = tbl.cell(0, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = head_fill
        put(cell, htxt, bold=True, color=INK, sz=hsize,
            align=(aligns[c] if aligns else PP_ALIGN.LEFT))
    for r_i, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = tbl.cell(r_i + 1, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if r_i % 2 == 0 else RGBColor(0xF8, 0xFA, 0xFC)
            col = (cell_colors or {}).get((r_i, c), INK)
            put(cell, val, bold=(first_bold and c == 0), color=col,
                align=(aligns[c] if aligns else PP_ALIGN.LEFT))
    return tbl


def _text_w(s, size):
    """全角=1.0em、半角=0.55em で概算した文字列幅(インチ)。"""
    em = size / 72.0
    n = sum(1.0 if ord(ch) > 0x2E7F else 0.55 for ch in s)
    return n * em


def chip_row(sl, x, y, items, h=0.28, gap=0.09, size=9.5):
    cx = x
    for label, color, fill in items:
        w = 0.22 + _text_w(label, size)
        badge(sl, cx, y, w, h, label, color=color, fill=fill, size=size)
        cx += w + gap
    return cx
